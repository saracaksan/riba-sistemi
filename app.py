# ==============================================================================
# SİSTEMİ HAZIRLAYAN: Sıraç AKSAN
# MEB RİBA Ortaokul Kapsamlı Yönetim, Takip ve Raporlama Otomasyonu
# ==============================================================================
from flask import Flask, render_template, jsonify, request, session, send_file, redirect
from supabase import create_client, Client
import os, io, csv, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Uygulama ve Klasör Ayarları (Logo için static klasörü tanımlandı)
sabit_klasor_yolu = os.path.dirname(os.path.abspath(__file__))
template_yonlendirici = os.path.join(sabit_klasor_yolu, 'templates')
static_yonlendirici = os.path.join(sabit_klasor_yolu, 'static')

app = Flask(__name__, template_folder=template_yonlendirici, static_folder=static_yonlendirici)
app.secret_key = "riba_2026_tam_entegrasyon_sirac_aksan"

# Supabase Veritabanı Bağlantısı
SUPABASE_URL = "https://pmlgahbdpzlhzlbvfpxf.supabase.co"
SUPABASE_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InBtbGdhaGJkcHpsaHpsYnZmcHhmIiwicm9sZSI6ImFub24iLCJpYXQiOjE3NzkyMjI4MzQsImV4cCI6MjA5NDc5ODgzNH0.6-OdE2chR29IJHKV1lCHYDkxE5HkkMtvMmXEkFYnbH0"
supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)

ILCE_USER = "Gazi"
ILCE_PASS = "Gazi.47"

# ================= 🌐 SAYFA YÖNLENDİRMELERİ =================
@app.route('/')
def ana_sayfa(): return render_template('index.html')

@app.route('/anket')
def anket_formu(): return render_template('anket.html')

@app.route('/panel/ilce-mem')
def panel_ilce(): 
    if session.get('rol') != 'ilce': return redirect('/')
    return render_template('panel_ilce.html')

@app.route('/panel/okul-idare')
def panel_okul():
    if session.get('rol') != 'okul': return redirect('/')
    return render_template('panel_okul.html')

@app.route('/panel/ogretmen')
def panel_ogretmen():
    if session.get('rol') != 'ogretmen': return redirect('/')
    return render_template('panel_ogretmen.html')

# ================= 🔐 KİMLİK DOĞRULAMA =================
@app.route('/api/personel-login', methods=['POST'])
def api_personel_login():
    veri = request.get_json()
    k_adi = str(veri.get('kullanici_adi', '')).strip()
    sifre = str(veri.get('sifre', '')).strip()
    
    if k_adi == ILCE_USER and sifre == ILCE_PASS:
        session.clear(); session['rol'] = 'ilce'
        return jsonify({"durum": "basarili", "yonlendir": "/panel/ilce-mem"})
    
    try:
        res = supabase.table("kullanicilar").select("*, okullar(okul_adi, kademe)").eq("kullanici_adi", k_adi).eq("sifre", sifre).execute()
        if res.data:
            kul = res.data[0]
            session.clear(); session['rol'] = kul['rol']; session['okul_id'] = kul['okul_id']; session['kullanici_adi'] = kul['kullanici_adi']
            if kul['rol'] == 'okul': return jsonify({"durum": "basarili", "yonlendir": "/panel/okul-idare"})
            if kul['rol'] == 'ogretmen': return jsonify({"durum": "basarili", "yonlendir": "/panel/ogretmen"})
        return jsonify({"durum": "hata", "mesaj": "Kullanıcı adı veya şifre hatalı!"}), 401
    except Exception as e: return jsonify({"durum": "hata", "mesaj": str(e)}), 500

@app.route('/api/ogrenci-giris', methods=['POST'])
def api_ogrenci_giris():
    veri = request.get_json()
    try:
        res = supabase.table("ogrenciler").select("id, ad, soyad, siniflar(sinif_adi)").eq("okul_id", veri.get('okul_id')).eq("okul_no", veri.get('okul_no')).execute()
        if res.data:
            ogr = res.data[0]
            tam_ad = f"{ogr.get('ad','')} {ogr.get('soyad','')}".strip()
            return jsonify({"durum": "basarili", "id": ogr['id'], "ad_soyad": tam_ad, "sinif": ogr['siniflar']['sinif_adi']})
        return jsonify({"durum": "hata", "mesaj": "Kayıt bulunamadı!"}), 404
    except: return jsonify({"durum": "hata"}), 500

@app.route('/api/okullari-getir')
@app.route('/api/ilce/okullari-getir')
def api_okullari_listele():
    try:
        res = supabase.table("okullar").select("id, okul_adi, kademe, kurum_kodu, kullanicilar(sifre)").execute()
        return jsonify([{"id": o['id'], "okul_adi": o['okul_adi'], "kademe": o['kademe'], "kurum_kodu": o['kurum_kodu'], "sifre": o['kullanicilar'][0]['sifre'] if o.get('kullanicilar') else ""} for o in res.data])
    except: return jsonify([])

# ================= 🏛️ İLÇE MEM: OKUL İŞLEMLERİ =================
@app.route('/api/ilce/okul-ekle', methods=['POST'])
def api_okul_ekle():
    veri = request.get_json()
    try:
        y_okul = supabase.table("okullar").insert({"okul_adi": veri['okul_adi'], "kademe": veri['kademe'], "kurum_kodu": str(veri['kurum_kodu'])}).execute()
        supabase.table("kullanicilar").insert({"kullanici_adi": str(veri['kurum_kodu']), "sifre": veri['sifre'], "rol": "okul", "okul_id": y_okul.data[0]['id']}).execute()
        return jsonify({"durum": "basarili", "mesaj": "Okul eklendi!"})
    except: return jsonify({"durum": "hata"})

@app.route('/api/ilce/okul-sil/<int:id>', methods=['DELETE'])
def api_okul_sil(id):
    try: supabase.table("okullar").delete().eq("id", id).execute(); return jsonify({"durum": "basarili"})
    except: return jsonify({"durum": "hata"})

# ================= 👨‍🏫 OKUL İDARESİ: ÖĞRETMEN VE ŞUBE YÖNETİMİ =================
@app.route('/api/okul/ogretmen-ekle', methods=['POST'])
def api_ogretmen_ekle():
    veri = request.get_json(); okul_id = session.get('okul_id')
    try:
        supabase.table("kullanicilar").insert({
            "kullanici_adi": veri['ad_soyad'], "sifre": veri['sifre'], "rol": "ogretmen", 
            "okul_id": okul_id, "ad_soyad": veri['ad_soyad'], "atanan_siniflar": veri.get('atanan_siniflar', 'TÜMÜ')
        }).execute()
        return jsonify({"durum": "basarili", "mesaj": "Öğretmen eklendi ve şubeleri atandı!"})
    except Exception as e: return jsonify({"durum": "hata", "mesaj": str(e)})

@app.route('/api/okul/ogretmenleri-getir')
def api_ogretmenleri_getir():
    try: return jsonify(supabase.table("kullanicilar").select("id, ad_soyad, sifre, atanan_siniflar, anket_doldurdu").eq("okul_id", session.get('okul_id')).eq("rol", "ogretmen").execute().data)
    except: return jsonify([])

@app.route('/api/okul/ogretmen-sil/<int:id>', methods=['DELETE'])
def api_ogretmen_sil(id):
    try: supabase.table("kullanicilar").delete().eq("id", id).execute(); return jsonify({"durum": "basarili"})
    except: return jsonify({"durum": "hata"})

# ================= 📋 OKUL İDARESİ: ÖĞRENCİ YÖNETİMİ & TOPLU DÜZELTME =================
@app.route('/api/okul/ogrencileri-getir')
def api_ogrencileri_getir():
    try: return jsonify(supabase.table("ogrenciler").select("id, okul_no, ad, soyad, cinsiyet, siniflar(sinif_adi)").eq("okul_id", session.get('okul_id')).order("okul_no").execute().data)
    except: return jsonify([])

@app.route('/api/okul/ogrenci-ekle-manuel', methods=['POST'])
def api_ogrenci_ekle_manuel():
    veri = request.get_json(); okul_id = session.get('okul_id')
    try:
        s_k = supabase.table("siniflar").select("id").eq("okul_id", okul_id).eq("sinif_adi", veri['sinif_adi']).execute()
        s_id = s_k.data[0]['id'] if s_k.data else supabase.table("siniflar").insert({"okul_id": okul_id, "sinif_adi": veri['sinif_adi']}).execute().data[0]['id']
        supabase.table("ogrenciler").insert({"okul_id": okul_id, "sinif_id": s_id, "okul_no": int(veri['okul_no']), "ad": veri['ad'], "soyad": veri['soyad'], "cinsiyet": veri.get('cinsiyet','Belirtilmemiş')}).execute()
        return jsonify({"durum": "basarili", "mesaj": "Öğrenci Eklendi!"})
    except: return jsonify({"durum": "hata"})

@app.route('/api/okul/ogrenci-guncelle/<int:id>', methods=['PUT'])
def api_ogrenci_guncelle(id):
    veri = request.get_json(); okul_id = session.get('okul_id')
    try:
        s_k = supabase.table("siniflar").select("id").eq("okul_id", okul_id).eq("sinif_adi", veri['sinif_adi']).execute()
        s_id = s_k.data[0]['id'] if s_k.data else supabase.table("siniflar").insert({"okul_id": okul_id, "sinif_adi": veri['sinif_adi']}).execute().data[0]['id']
        supabase.table("ogrenciler").update({"okul_no": int(veri['okul_no']), "ad": veri['ad'], "soyad": veri['soyad'], "cinsiyet": veri['cinsiyet'], "sinif_id": s_id}).eq("id", id).execute()
        return jsonify({"durum": "basarili", "mesaj": "Güncellendi!"})
    except: return jsonify({"durum": "hata"})

@app.route('/api/okul/ogrenci-sil/<int:id>', methods=['DELETE'])
def api_ogrenci_sil(id):
    try: supabase.table("ogrenciler").delete().eq("id", id).execute(); return jsonify({"durum": "basarili"})
    except: return jsonify({"durum": "hata"})

@app.route('/api/okul/ogrenci-toplu-sinif-guncelle', methods=['POST'])
def api_ogrenci_toplu_sinif_guncelle():
    veri = request.get_json(); idler = veri.get('idler', []); yeni_sinif_adi = veri.get('yeni_sinif_adi'); okul_id = session.get('okul_id')
    if not idler or not yeni_sinif_adi: return jsonify({"durum": "hata", "mesaj": "Eksik bilgi gönderildi."})
    try:
        s_k = supabase.table("siniflar").select("id").eq("okul_id", okul_id).eq("sinif_adi", yeni_sinif_adi).execute()
        s_id = s_k.data[0]['id'] if s_k.data else supabase.table("siniflar").insert({"okul_id": okul_id, "sinif_adi": yeni_sinif_adi}).execute().data[0]['id']
        supabase.table("ogrenciler").update({"sinif_id": s_id}).in_("id", idler).execute()
        return jsonify({"durum": "basarili", "mesaj": f"{len(idler)} öğrencinin sınıfı ışık hızında '{yeni_sinif_adi}' olarak güncellendi!"})
    except Exception as e: return jsonify({"durum": "hata", "mesaj": str(e)})

@app.route('/api/okul/ogrenci-toplu-sil', methods=['POST'])
def api_ogrenci_toplu_sil():
    veri = request.get_json(); idler = veri.get('idler', [])
    if not idler: return jsonify({"durum": "hata", "mesaj": "Silinecek öğrenci bulunamadı."})
    try:
        supabase.table("ogrenciler").delete().in_("id", idler).execute()
        return jsonify({"durum": "basarili", "mesaj": f"{len(idler)} öğrenci silindi!"})
    except Exception as e: return jsonify({"durum": "hata", "mesaj": str(e)})

# ================= ⚡ EXCEL ŞABLON İNDİRME VE YÜKLEME =================
@app.route('/api/okul/sablon-indir')
def api_okul_sablon_indir():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Ogrenciler"
    ws.append(["Okul No", "Adı", "Soyadı", "Cinsiyeti", "Sınıfı"])
    ws.append([101, "Örnek Ad", "Örnek Soyad", "Erkek", "5/A"])
    for col in range(1, 6): 
        ws.cell(row=1, column=col).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=1, column=col).fill = PatternFill(start_color="1E3A8A", fill_type="solid")
    out = io.BytesIO(); wb.save(out); out.seek(0)
    return send_file(out, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name="Okul_Veri_Yukleme_Sablonu.xlsx")

@app.route('/api/sablon-liste-yukle', methods=['POST'])
def api_sablon_yukle():
    if 'file' not in request.files: return jsonify({"durum": "hata", "mesaj": "Dosya yok"}), 400
    file = request.files['file']; okul_id = session.get('okul_id'); filename = file.filename.lower(); eklenen = 0
    try:
        if filename.endswith('.csv'):
            stream = io.StringIO(file.stream.read().decode("UTF8"), newline=None); csv_reader = csv.reader(stream, delimiter=',')
            next(csv_reader, None)
            for row in csv_reader:
                if len(row) < 5: continue
                no, ad, soyad, cinsiyet, sinif = [str(e).strip() for e in row[:5]]
                if not no.isdigit(): continue
                s_k = supabase.table("siniflar").select("id").eq("okul_id", okul_id).eq("sinif_adi", sinif).execute()
                s_id = s_k.data[0]['id'] if s_k.data else supabase.table("siniflar").insert({"okul_id": okul_id, "sinif_adi": sinif}).execute().data[0]['id']
                try: supabase.table("ogrenciler").insert({"okul_id": okul_id, "sinif_id": s_id, "okul_no": int(no), "ad": ad, "soyad": soyad, "cinsiyet": cinsiyet}).execute(); eklenen += 1
                except: pass
        elif filename.endswith(('.xls', '.xlsx')):
            wb = openpyxl.load_workbook(file); ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[0] or not row[1] or not row[4]: continue
                no = str(row[0]).strip(); ad = str(row[1]).strip(); soyad = str(row[2]).strip() if row[2] else ""; cinsiyet = str(row[3]).strip() if row[3] else "Belirtilmemiş"; sinif = str(row[4]).strip()
                if not no.isdigit(): continue
                s_k = supabase.table("siniflar").select("id").eq("okul_id", okul_id).eq("sinif_adi", sinif).execute()
                s_id = s_k.data[0]['id'] if s_k.data else supabase.table("siniflar").insert({"okul_id": okul_id, "sinif_adi": sinif}).execute().data[0]['id']
                try: supabase.table("ogrenciler").insert({"okul_id": okul_id, "sinif_id": s_id, "okul_no": int(no), "ad": ad, "soyad": soyad, "cinsiyet": cinsiyet}).execute(); eklenen += 1
                except: pass
        return jsonify({"durum": "basarili", "mesaj": f"{eklenen} öğrenci başarıyla aktarıldı!"})
    except Exception as e: return jsonify({"durum": "hata", "mesaj": str(e)}), 500

# ================= 📝 ANKET KAYDETME MOTORU =================
@app.route('/api/anket-kaydet', methods=['POST'])
def api_anket_kaydet():
    veri = request.get_json()
    try:
        supabase.table("riba_cevaplari").insert({
            "ogrenci_id": veri.get('ogrenci_id'), "dolduran_tipi": veri['dolduran_tipi'], "cevaplar": veri['cevaplar']
        }).execute()
        if veri['dolduran_tipi'] == 'Ogretmen':
            k_adi = session.get('kullanici_adi')
            if k_adi: supabase.table("kullanicilar").update({"anket_doldurdu": True}).eq("kullanici_adi", k_adi).execute()
        return jsonify({"durum": "basarili"})
    except Exception as e: return jsonify({"durum": "hata", "mesaj": str(e)})

# ================= 👨‍🏫 ÖĞRETMEN TAKİP MOTORU =================
@app.route('/api/ogretmen/sinif-takip')
def api_ogretmen_takip():
    okul_id = session.get('okul_id'); k_adi = session.get('kullanici_adi')
    try:
        kul = supabase.table("kullanicilar").select("atanan_siniflar").eq("kullanici_adi", k_adi).execute().data[0]
        atanan_str = kul.get('atanan_siniflar', 'TÜMÜ')
        atanan_liste = [s.strip() for s in atanan_str.split(',')] if atanan_str != 'TÜMÜ' else []
        
        ogrenciler = supabase.table("ogrenciler").select("id, okul_no, ad, soyad, siniflar(sinif_adi)").eq("okul_id", okul_id).execute().data
        if atanan_str != 'TÜMÜ': ogrenciler = [o for o in ogrenciler if o['siniflar']['sinif_adi'] in atanan_liste]
            
        cevaplar = supabase.table("riba_cevaplari").select("ogrenci_id, dolduran_tipi").execute().data
        o_cev = {c['ogrenci_id']: True for c in cevaplar if c['dolduran_tipi'] == 'Ogrenci'}
        v_cev = {c['ogrenci_id']: True for c in cevaplar if c['dolduran_tipi'] == 'Veli'}
        
        liste = [{"no": o['okul_no'], "ad_soyad": f"{o['ad']} {o['soyad']}", "sinif": o['siniflar']['sinif_adi'], "ogrenci_durum": "Tamamlandı" if o['id'] in o_cev else "Eksik", "veli_durum": "Tamamlandı" if o['id'] in v_cev else "Eksik"} for o in ogrenciler]
        return jsonify({"ogrenciler": liste, "atanan_siniflar": atanan_str})
    except Exception as e: return jsonify({"durum": "hata", "mesaj": str(e)}), 500

@app.route('/api/takip/katilim-durumu')
def api_katilim_durumu():
    rol = session.get('rol'); okul_id = request.args.get('okul_id') if rol == 'ilce' else session.get('okul_id')
    try:
        q = supabase.table("ogrenciler").select("id, okul_no, ad, soyad, sinif_id, siniflar(sinif_adi)")
        if okul_id: q = q.eq("okul_id", okul_id)
        ogrenciler = q.execute().data
        cevaplar = supabase.table("riba_cevaplari").select("ogrenci_id, dolduran_tipi").execute().data
        
        o_cev = {c['ogrenci_id']: True for c in cevaplar if c['dolduran_tipi'] == 'Ogrenci'}
        v_cev = {c['ogrenci_id']: True for c in cevaplar if c['dolduran_tipi'] == 'Veli'}
        s_mevcut = {}; s_katilim = {"ogrenci": {}, "veli": {}}
        for o in ogrenciler:
            s_adi = o['siniflar']['sinif_adi']
            s_mevcut[s_adi] = s_mevcut.get(s_adi, 0) + 1
            if o['id'] in o_cev: s_katilim["ogrenci"][s_adi] = s_katilim["ogrenci"].get(s_adi, 0) + 1
            if o['id'] in v_cev: s_katilim["veli"][s_adi] = s_katilim["veli"].get(s_adi, 0) + 1

        eksik = [{"sinif_adi": s, "mevcut": m, "ogrenci_katilim": s_katilim["ogrenci"].get(s,0), "veli_katilim": s_katilim["veli"].get(s,0)} for s, m in s_mevcut.items() if s_katilim["ogrenci"].get(s,0) < m or s_katilim["veli"].get(s,0) < m]
        return jsonify({"eksik_siniflar": eksik, "toplam_mevcut": len(ogrenciler), "toplam_ogrenci_katilim": len(o_cev), "toplam_veli_katilim": len(v_cev)})
    except: return jsonify({"durum": "hata"}), 500

# ================= 📈 OKUL/SINIF/İLÇE GENEL RAPORLAMA BEYNİ =================
def analiz_hesapla(okul_id, sinif_adi):
    hedefler_map = {
        "Okula ve Çevreye Uyum/Okul Kuralları": ["M1"], "Hedef Belirleme": ["M2"], "Yetenekleri Tanıma": ["M3"],
        "İlgileri Keşfetme": ["M4"], "Mesleki Benlik": ["M5"], "Zaman Yönetimi": ["M6"],
        "Öz Düzenlemeli Öğrenme (Ders Çalışma Becerileri)": ["M7"], "Öz Disiplin Geliştirme (Sorumluluk)": ["M8"],
        "Dikkat Geliştirme Çalışmaları": ["M9"], "Sınav Kaygısı": ["M10"], "Motivasyon": ["M11"], "Devamsızlığı Önleme": ["M12"],
        "Karar Verme Becerisi": ["M13"], "Problem Çözme Becerileri": ["M14"], "Üst Öğrenim Kurumlarının Tanıtılması": ["M15"],
        "Üst Öğrenime Geçiş Sınavları": ["M16"], "Özgüven Geliştirme": ["M17"], "Duygu Farkındalığı": ["M18"],
        "Duygu Düzenleme": ["M19"], "Öfke Yönetimi": ["M20"], "Hak ve Sorumluluklarını Bilme": ["M21"], "Atılganlık": ["M22"],
        "İletişim Becerileri": ["M23"], "Kendini İfade Etme": ["M24"], "Sosyal Beceriler (Arkadaşlık)": ["M25"],
        "Bireysel Farklılıklara Saygı": ["M26"], "Sınır Koyma (HAYIR diyebilmek)": ["M27"], "Çatışma Çözme Becerileri": ["M28"],
        "Gelişim Dönemi Özellikleri (Ergenlik)": ["M29"], "Sağlıklı Yaşam": ["M30"], "İhmal ve İstismardan Korunma": ["M31"],
        "Yaşam Becerileri (Riskli Davranışlar)": ["M32"], "Akran Zorbalığı": ["M33"], "Bilinçli Teknoloji Kullanımı": ["M34"],
        "Bağımlılıkla Mücadele": ["M35"], "Psikolojik Sağlamlık": ["M36"]
    }
    
    frekanslar = {isim: {"Ogrenci": 0, "Veli": 0, "Ogretmen": 0} for isim in hedefler_map.keys()}
    sorgu = supabase.table("riba_cevaplari").select("*, ogrenciler(okul_id, siniflar(sinif_adi))")
    cevaplar = sorgu.execute().data
    
    toplam_ogr = 0; toplam_veli = 0; toplam_ogrt = 0
    for c in cevaplar:
        ogr_baglanti = c.get('ogrenciler')
        if not ogr_baglanti and c['dolduran_tipi'] != 'Ogretmen': continue
        if ogr_baglanti:
            if okul_id and str(ogr_baglanti['okul_id']) != str(okul_id): continue
            if sinif_adi and ogr_baglanti['siniflar']['sinif_adi'] != sinif_adi: continue
        
        tip = c['dolduran_tipi']; c_json = c.get('cevaplar', {})
        if tip == "Ogrenci": toplam_ogr += 1
        elif tip == "Veli": toplam_veli += 1
        elif tip == "Ogretmen": toplam_ogrt += 1

        for hedef_adi, kodlar in hedefler_map.items():
            for kod in kodlar:
                if kod in c_json: frekanslar[hedef_adi][tip] += c_json[kod].get("A", 0)

    rapor_matrisi = []
    for hedef_adi, frek in frekanslar.items():
        f_o = frek["Ogrenci"]; f_v = frek["Veli"]; f_og = frek["Ogretmen"]
        asp = round((f_o * 0.5) + (f_v * 0.3) + (f_og * 0.2), 2)
        rapor_matrisi.append({"hedef_adi": hedef_adi, "ogrenci_f": f_o, "veli_f": f_v, "ogretmen_f": f_og, "asp": asp})
        
    rapor_matrisi.sort(key=lambda x: x['asp'], reverse=True)
    return rapor_matrisi, {"ogrenci": toplam_ogr, "veli": toplam_veli, "ogretmen": toplam_ogrt}

@app.route('/api/rapor/analiz-veri')
def api_rapor_analiz_veri():
    rol = session.get('rol')
    okul_id = request.args.get('okul_id') if rol == 'ilce' else session.get('okul_id')
    sinif_adi = request.args.get('sinif_adi')
    rapor, _ = analiz_hesapla(okul_id, sinif_adi)
    return jsonify(rapor)

@app.route('/api/rapor/sinif-listesi')
def api_rapor_sinif_listesi():
    okul_id = session.get('okul_id')
    if not okul_id: return jsonify([])
    res = supabase.table("siniflar").select("sinif_adi").eq("okul_id", okul_id).execute().data
    
    # MÜKERRER SINIFLARI TEKE DÜŞÜRME FİLTRESİ
    benzersiz_siniflar = sorted(list(set([s['sinif_adi'] for s in res if s['sinif_adi']])))
    return jsonify(benzersiz_siniflar)

@app.route('/api/rapor/excel-indir')
def api_excel_indir():
    rol = session.get('rol')
    okul_id = request.args.get('okul_id') if rol == 'ilce' else session.get('okul_id')
    sinif_adi = request.args.get('sinif_adi')
    
    rapor, katilim = analiz_hesapla(okul_id, sinif_adi)
    baslik = f"{sinif_adi} Sınıfı RİBA Sonuç Raporu" if sinif_adi else "Okul Geneli RİBA Sonuç Raporu"
    if not okul_id: baslik = "İlçe Geneli Toplu RİBA Sonuç Raporu"

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "RİBA Raporu"
    ws.append([baslik])
    ws.append([f"Katılım -> Öğrenci: {katilim['ogrenci']}, Veli: {katilim['veli']}, Öğretmen: {katilim['ogretmen']}"])
    ws.append([])
    
    headers = ["Sıra", "Rehberlik Hedefi", "Öğrenci (f)", "Veli (f)", "Öğretmen (f)", "ASP Puanı"]
    ws.append(headers)
    
    header_fill = PatternFill(start_color="1E3A8A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    for col in range(1, 7):
        ws.cell(row=4, column=col).fill = header_fill
        ws.cell(row=4, column=col).font = header_font
        ws.cell(row=4, column=col).alignment = Alignment(horizontal="center")

    for i, r in enumerate(rapor):
        ws.append([i+1, r['hedef_adi'], r['ogrenci_f'], r['veli_f'], r['ogretmen_f'], r['asp']])
    
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=1, max_col=6):
        for cell in row:
            cell.border = border
            if cell.column != 2: cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions['A'].width = 10; ws.column_dimensions['B'].width = 50
    ws.column_dimensions['C'].width = 15; ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15; ws.column_dimensions['F'].width = 15

    out = io.BytesIO(); wb.save(out); out.seek(0)
    return send_file(out, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name=f"RIBA_Sonuc_{'Sinif' if sinif_adi else 'Okul'}.xlsx")

if __name__ == '__main__':
    print("Sistemi Hazırlayan: Sıraç AKSAN")
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 5000)), debug=True)